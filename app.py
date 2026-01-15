import pandas as pd
import numpy as np
import pyodbc
import warnings
import os
import shutil
import tempfile
import openpyxl
import streamlit as st
import time

# --- 1. CONFIGURACION DE LA APP ---
st.set_page_config(
    page_title="Monitor de Stock Brogas",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="collapsed"
)

warnings.filterwarnings('ignore')

# --- CONSTANTES ---
PATH_EXCEL_ORIGEN = r"O:\TALLERES 2\Proyectado de 6 meses.xlsx"
FECHA_FILTRO_BROGAS = '2024-12-01'
FECHA_FILTRO_ML = '2025-09-01'

# --- 2. GESTIÓN DE CACHÉ Y CONEXIONES ---

def conectar_odbc(dsn):
    try:
        return pyodbc.connect(f"DSN={dsn};Uid=SYSDBA;Pwd=masterkey", timeout=10)
    except Exception as e:
        st.error(f"❌ Error conectando a {dsn}: {e}")
        return None

@st.cache_data(ttl=3600, show_spinner="Leyendo Excel Proyectado...")
def get_proyectado_optimizado():
    if not os.path.exists(PATH_EXCEL_ORIGEN):
        st.error(f"No se encuentra el archivo: {PATH_EXCEL_ORIGEN}")
        return pd.DataFrame()

    with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
        try:
            shutil.copy2(PATH_EXCEL_ORIGEN, tmp.name)
            path_lectura = tmp.name
        except PermissionError:
            path_lectura = PATH_EXCEL_ORIGEN

    try:
        wb = openpyxl.load_workbook(path_lectura, data_only=True, read_only=False)
        res_data = None
        for sheet in wb.worksheets:
            if "PROYECTADO_2" in sheet.tables:
                tbl = sheet.tables["PROYECTADO_2"]
                rango = tbl.ref
                # CORRECCIÓN PYLANCE: Extracción limpia de valores
                data = []
                for row in sheet[rango]:
                    data.append([cell.value for cell in row]) # type: ignore
                
                if len(data) > 1:
                    res_data = pd.DataFrame(data[1:], columns=data[0])
                break
        wb.close()
        
        try: os.remove(path_lectura)
        except: pass

        if res_data is None: return pd.DataFrame()

        df = res_data
        df.columns = df.columns.astype(str).str.strip()
        meses = ['MES2', 'MES3', 'MES4']
        for m in meses:
            if m in df.columns:
                df[m] = pd.to_numeric(df[m], errors='coerce').fillna(0)
            else:
                df[m] = 0
        
        df['PEDIDO_PROYECTADO'] = df[meses].sum(axis=1)
        col_cod = 'Codigo' if 'Codigo' in df.columns else df.columns[0]
        col_des = 'Descripción' if 'Descripción' in df.columns else (df.columns[1] if len(df.columns) > 1 else 'Descripción')
        
        df = df[[col_cod, col_des, 'PEDIDO_PROYECTADO']].rename(columns={col_cod: 'CODIGOPARTICULAR', col_des: 'DESCRIPCION'})
        df['CODIGOPARTICULAR'] = df['CODIGOPARTICULAR'].astype(str).str.strip().str.upper()
        return df
    except Exception as e:
        st.error(f"Error procesando Excel: {e}")
        return pd.DataFrame()

@st.cache_data(ttl=600, show_spinner="Consultando Base de Datos...")
def get_datos_sql():
    conn_brogas = conectar_odbc("BROGAS")
    conn_ml = conectar_odbc("BROGASML")
    if not conn_brogas: return None, None, None, None, None

    try:
        # CORRECCIÓN PYLANCE: se añade # type: ignore para las conexiones de pyodbc en read_sql
        q_art = "SELECT A.CODIGOPARTICULAR, A.DESCRIPCION, SUM(C.STOCKACTUAL) as STOCK FROM ARTICULOS A LEFT JOIN CASILLEROS C ON A.CODIGOARTICULO = C.CODIGOARTICULO LEFT JOIN DEPOSITOS D ON C.CODIGODEPOSITO = D.CODIGODEPOSITO WHERE D.DESCRIPCION NOT IN ('COMPRAS NC','ALUCOLOR','ECOMMERCE_FULL_BRO','ECOMMERCE_FULL_1','CONTROL DE CALIDAD', 'SALDOS','ECOMMERCE_FACTURACIÓN', 'ECOMMERCE_STOCK', 'SCRAP', 'SERVICIO TECNICO', 'SHOWROOM', 'M. NO CONFORMES') GROUP BY A.CODIGOPARTICULAR, A.DESCRIPCION"
        df_art = pd.read_sql(q_art, conn_brogas) # type: ignore

        q_ventas = f"SELECT CODIGOPARTICULAR, SUM(CANTIDAD - CANTIDADREMITIDA) as PENDIENTES_VENTAS FROM CUERPOCOMPROBANTES WHERE FECHAMODIFICACION > '{FECHA_FILTRO_BROGAS}' AND TIPOCOMPROBANTE IN ('FA', 'FB', 'FCA', 'FE') AND (CANTIDAD - CANTIDADREMITIDA) > 0 GROUP BY CODIGOPARTICULAR"
        df_ventas = pd.read_sql(q_ventas, conn_brogas) # type: ignore

        q_pedidos = "SELECT CP.CODIGOPARTICULAR, SUM(CP.CANTIDAD) as PEDIDOS_NUEVOS FROM CUERPOPEDIDOS CP INNER JOIN CABEZAPEDIDOS CB ON CP.NUMEROCOMPROBANTE = CB.NUMEROCOMPROBANTE AND CP.TIPOCOMPROBANTE = CB.TIPOCOMPROBANTE INNER JOIN DEPOSITOS D ON CP.CODIGODEPOSITO = D.CODIGODEPOSITO WHERE CB.ANULADA = 0 AND CP.CANTIDADCANCELADA = 0 AND CP.CANTIDADREMITIDA = 0 AND CP.CANTIDADPREPARADA = 0 AND D.DESCRIPCION IN ('EXPEDICION', 'FIZBAY') GROUP BY CP.CODIGOPARTICULAR"
        df_pedidos = pd.read_sql(q_pedidos, conn_brogas) # type: ignore

        q_op = "SELECT A.CODIGOPARTICULAR, SUM(CP.CANTIDAD) as CANTIDAD_TOTAL_OP, SUM(CP.CANTIDAD - COALESCE(ENTREGAS.TOTAL_ENTREGADO, 0)) as EN_PRODUCCION FROM PRODCABEZAORDEN H INNER JOIN PRODCUERPOORDEN CP ON H.CODIGOORDEN = CP.CODIGOORDEN INNER JOIN ARTICULOS A ON CP.CODIGOARTICULO = A.CODIGOARTICULO LEFT JOIN (SELECT CODIGOORDEN, CODIGOARTICULO, SUM(CANTIDAD) as TOTAL_ENTREGADO FROM PRODDETALLEFINALIZACIONORDEN GROUP BY CODIGOORDEN, CODIGOARTICULO) ENTREGAS ON CP.CODIGOORDEN = ENTREGAS.CODIGOORDEN AND CP.CODIGOARTICULO = ENTREGAS.CODIGOARTICULO LEFT JOIN ESTADOSORDENPRODUCCION E ON H.CODIGOESTADOOP = E.CODIGOESTADOOP WHERE H.ANULADA = 0 AND COALESCE(E.DESCRIPCION, '') <> 'TERMINADO' AND (CP.CANTIDAD - COALESCE(ENTREGAS.TOTAL_ENTREGADO, 0)) > 0 GROUP BY A.CODIGOPARTICULAR"
        df_op = pd.read_sql(q_op, conn_brogas) # type: ignore
    finally:
        conn_brogas.close()

    df_ml = pd.DataFrame()
    if conn_ml:
        try:
            q_ml = f"SELECT CODIGOPARTICULAR, SUM(CANTIDAD - CANTIDADREMITIDA) as PENDIENTE_ML FROM CUERPOCOMPROBANTES WHERE FECHAMODIFICACION > '{FECHA_FILTRO_ML}' GROUP BY CODIGOPARTICULAR"
            df_ml = pd.read_sql(q_ml, conn_ml) # type: ignore
        except: pass
        finally: conn_ml.close()

    return df_art, df_ventas, df_pedidos, df_op, df_ml

# --- 3. LÓGICA DE CONSOLIDACIÓN ---

def procesar_datos_consolidado():
    df_proy = get_proyectado_optimizado()
    if df_proy.empty: return pd.DataFrame(), pd.DataFrame(), pd.DataFrame()

    df_art, df_ventas, df_pedidos, df_op, df_ml = get_datos_sql()
    if df_art is None: return df_proy, pd.DataFrame(), pd.DataFrame()

    for d in [df_art, df_ventas, df_pedidos, df_op, df_ml]:
        if d is not None and not d.empty: d.columns = d.columns.str.upper()

    final = df_proy.merge(df_art, on='CODIGOPARTICULAR', how='left', suffixes=('_EXCEL', '_SQL'))
    final['DESCRIPCION'] = final['DESCRIPCION_SQL'].fillna(final['DESCRIPCION_EXCEL'])
    final.drop(columns=['DESCRIPCION_SQL', 'DESCRIPCION_EXCEL'], inplace=True, errors='ignore')

    if df_ventas is not None: final = final.merge(df_ventas, on='CODIGOPARTICULAR', how='left')
    if df_pedidos is not None: final = final.merge(df_pedidos, on='CODIGOPARTICULAR', how='left')
    if df_ml is not None and not df_ml.empty: final = final.merge(df_ml, on='CODIGOPARTICULAR', how='left')
    if df_op is not None: final = final.merge(df_op, on='CODIGOPARTICULAR', how='left')

    final = final.fillna(0)
    final['PENDIENTE_TOTAL'] = final.get('PENDIENTES_VENTAS', 0) + final.get('PEDIDOS_NUEVOS', 0) + final.get('PENDIENTE_ML', 0)
    final['STOCK_NETO'] = final['STOCK'] - final['PENDIENTE_TOTAL']
    final['COBERTURA_MESES'] = np.where(final['PEDIDO_PROYECTADO'] > 0, (final['STOCK_NETO'] / final['PEDIDO_PROYECTADO']) * 3, 999)
    final.loc[(final['STOCK_NETO'] <= 0), 'COBERTURA_MESES'] = 0

    cols_order = ['CODIGOPARTICULAR', 'DESCRIPCION', 'PEDIDO_PROYECTADO', 'STOCK', 'PENDIENTE_TOTAL', 'STOCK_NETO', 'COBERTURA_MESES', 'EN_PRODUCCION']
    for c in cols_order:
        if c not in final.columns: final[c] = 0

    return final[cols_order], df_art, df_op

# --- 4. INTERFAZ VISUAL ---

def main():
    st.title("🏭 Monitor de Stock e Inventario")
    
    col1, col2 = st.columns([4, 1])
    with col1: st.caption(f"**Origen:** {PATH_EXCEL_ORIGEN}")
    with col2:
        if st.button("🔄 Actualizar", type="primary"):
            st.cache_data.clear()
            st.rerun()

    with st.spinner("Procesando..."):
        df_final, df_stock_bruto, df_prod_bruto = procesar_datos_consolidado()

    if df_final.empty:
        st.warning("⚠️ Sin datos.")
        return

    if df_prod_bruto is not None and not df_prod_bruto.empty:
        df_prod_bruto.columns = df_prod_bruto.columns.str.upper()

    codigos_excel = df_final['CODIGOPARTICULAR'].unique()
    tab1, tab2, tab3 = st.tabs(["🚀 Cobertura", "📦 Maestro Stock", "🛠️ Producción"])

    def formatear_y_mostrar(df_in):
        # ORDENAR ASCENDENTE Y QUITAR INDICE
        df_in = df_in.sort_values('CODIGOPARTICULAR', ascending=True)
        numeric_cols = df_in.select_dtypes(include=[np.number]).columns
        format_dict = {col: "{:,.0f}" for col in numeric_cols}
        st.dataframe(df_in.style.format(format_dict), use_container_width=True, hide_index=True)

    with tab1:
        def style_cobertura(val):
            color = "#ffffff" 
            if val <= 1.0: color = '#ff4b4b'
            elif val <= 2.0: color = '#ffa421'
            elif val > 2.0 and val < 100: color = '#21c354'
            return f'background-color: {color}; color: black'

        if df_prod_bruto is not None and not df_prod_bruto.empty:
            df_prod_temp = df_prod_bruto[['CODIGOPARTICULAR', 'EN_PRODUCCION']].copy()
            df_prod_temp = df_prod_temp.rename(columns={'EN_PRODUCCION': 'SALDO PENDIENTE'})
            df_mostrar = df_final.merge(df_prod_temp, on='CODIGOPARTICULAR', how='left')
        else:
            df_mostrar = df_final.copy()
            df_mostrar['SALDO PENDIENTE'] = 0

        df_mostrar['SALDO PENDIENTE'] = df_mostrar['SALDO PENDIENTE'].fillna(0)
        
        # ORDENAR ASCENDENTE
        df_mostrar = df_mostrar.sort_values('CODIGOPARTICULAR', ascending=True)

        cols_numericas = ['STOCK', 'STOCK_NETO', 'PEDIDO_PROYECTADO', 'PENDIENTE_TOTAL', 'EN_PRODUCCION', 'SALDO PENDIENTE']
        cols_a_formatear = [c for c in cols_numericas if c in df_mostrar.columns]

        st.dataframe(
            df_mostrar.style.map(style_cobertura, subset=['COBERTURA_MESES', 'STOCK_NETO'])
            .format("{:,.0f}", subset=cols_a_formatear)
            .format("{:,.2f}", subset=['COBERTURA_MESES'] if 'COBERTURA_MESES' in df_mostrar.columns else []),
            use_container_width=True,
            height=700,
            hide_index=True # QUITA LA PRIMERA COLUMNA
        )
        
        fecha = time.strftime("%Y%m%d_%H%M")
        st.download_button("📥 Descargar CSV", df_mostrar.to_csv(index=False).encode('utf-8'), f"Stock_{fecha}.csv", "text/csv")

    with tab2:
        st.subheader("Maestro Artículos (Filtrado por Excel)")
        if df_stock_bruto is not None and not df_stock_bruto.empty:
            df_stock_filtrado = df_stock_bruto[df_stock_bruto['CODIGOPARTICULAR'].isin(codigos_excel)].copy()
            formatear_y_mostrar(df_stock_filtrado)

    with tab3:
        st.subheader("Detalle de Producción (Filtrado por Excel)")
        if df_prod_bruto is not None and not df_prod_bruto.empty:
            df_prod_filtrado = df_prod_bruto[df_prod_bruto['CODIGOPARTICULAR'].isin(codigos_excel)].copy()
            df_descripciones = df_final[['CODIGOPARTICULAR', 'DESCRIPCION']].drop_duplicates()
            df_prod_filtrado = df_prod_filtrado.merge(df_descripciones, on='CODIGOPARTICULAR', how='left')
            df_prod_filtrado = df_prod_filtrado.rename(columns={'CANTIDAD_TOTAL_OP': 'CANT. TOTAL OP', 'EN_PRODUCCION': 'SALDO PENDIENTE'})
            
            cols_orden = ['CODIGOPARTICULAR', 'DESCRIPCION', 'CANT. TOTAL OP', 'SALDO PENDIENTE']
            df_prod_filtrado = df_prod_filtrado[[c for c in cols_orden if c in df_prod_filtrado.columns]]
            formatear_y_mostrar(df_prod_filtrado)

if __name__ == "__main__":
    main()